import streamlit as st
import fitz  # PyMuPDF
from docx import Document
from openpyxl import Workbook
from openai import OpenAI
import os
from dotenv import load_dotenv

#api groq
load_dotenv()
client = OpenAI(
    api_key=os.getenv("GROQ_API_KEY"),
    base_url="https://api.groq.com/openai/v1"
)


def extrair_texto_pdf(arquivo):
    texto = ""
    with fitz.open(stream=arquivo.read(), filetype='pdf') as doc:
        for pagina in doc:
            texto += pagina.get_text()
    return texto 


def gerar_resposta_groq(texto_pdf, modelo='llama3-70b-8192'):
    prompt = f'''
    A seguir está o conteúdo extraído do arquivo PDF:
    
    {texto_pdf}
    
    Organize essas informações de forma clara para gerar um arquivo Word e outro Excel, separando tópicos e tabelas.
    '''
    
    resposta = client.chat.completions.create(
        model=modelo,
        messages=[
            {"role": "system", "content": "Você é um assistente que organiza documentos extraídos de PDF."},
            {"role": "user", "content": prompt}
        ],
        temperature=0.3
    )
    return resposta.choices[0].message.content


def salvar_em_word(texto):
    doc = Document()
    for linha in texto.split("\n"):
        doc.add_paragraph(linha)
    caminho = "saida.docx"
    doc.save(caminho)
    return caminho



def salvar_em_excel(texto):
    wb = Workbook()
    ws = wb.active
    for linha in texto.strip().split("\n"):
        colunas = [col.strip() for col in linha.split(",")]
        ws.append(colunas)
    caminho = "saida.xlsx"
    wb.save(caminho)
    return caminho


# --- Interface Streamlit ---
st.set_page_config(page_title="GroqDoc - PDF para Word e Excel", layout="centered")
st.title("📄 GroqDoc: IA para transformar PDFs")
st.markdown("Envie um PDF que a IA vai organizar e gerar um Word e/ou Excel com os dados.")

arquivo_pdf = st.file_uploader("Selecione um arquivo PDF", type="pdf")

if arquivo_pdf:
    with st.spinner("🔍 Lendo e interpretando o PDF..."):
        texto_extraido = extrair_texto_pdf(arquivo_pdf)
        resposta = gerar_resposta_groq(texto_extraido)

    st.subheader("🧠 Texto interpretado:")
    st.text_area("Resultado da IA:", resposta, height=300)

    col1, col2 = st.columns(2)

    with col1:
        if st.button("📥 Baixar Word"):
            caminho_word = salvar_em_word(resposta)
            with open(caminho_word, "rb") as f:
                st.download_button("Download .docx", f, file_name="groqdoc.docx")

    with col2:
        if st.button("📥 Baixar Excel"):
            caminho_excel = salvar_em_excel(resposta)
            with open(caminho_excel, "rb") as f:
                st.download_button("Download .xlsx", f, file_name="groqdoc.xlsx")